"""Final two-level classification of the combined sheet (RESOLVED / UNRESOLVED).

    python -m verifier.classify_final            # final_combined/final_sorted.csv -> final_classified.*

RESOLVED:   VERIFIED_ACTIVE, REPLACED_ACTIVE, COMPANY_CLOSED, COMPANY_MERGED, NOT_FOUND_IN_SEARCH
UNRESOLVED: BLOCKED, SUGGESTED_DOMAIN_BUT_LOW_EVIDENCE

NEEDS_REVIEW rows with a suggested domain become REPLACED_ACTIVE (the suggestion is the final URL;
the accuracy audit found these ~75% correct). NEEDS_REVIEW rows without one become
SUGGESTED_DOMAIN_BUT_LOW_EVIDENCE: rows that could not be confidently mapped; their old provided
URL stays in "Old URL Provided". The input sheet is never modified.
"""
from __future__ import annotations

import argparse
import csv
from collections import Counter
from pathlib import Path

from .master_resolver import _save_workbook, excel_cell
from .normalization import normalize_domain

ROOT = Path(__file__).resolve().parent.parent
RESOLVED, UNRESOLVED = "RESOLVED", "UNRESOLVED"
LOW_EVIDENCE = "SUGGESTED_DOMAIN_BUT_LOW_EVIDENCE"
ORDER = [(RESOLVED, "VERIFIED_ACTIVE"), (RESOLVED, "REPLACED_ACTIVE"), (RESOLVED, "COMPANY_CLOSED"),
         (RESOLVED, "COMPANY_MERGED"), (RESOLVED, "NOT_FOUND_IN_SEARCH"),
         (UNRESOLVED, "BLOCKED"), (UNRESOLVED, LOW_EVIDENCE)]
DIRECT = {"VERIFIED_ACTIVE": (RESOLVED, "VERIFIED_ACTIVE"), "REPLACED_ACTIVE": (RESOLVED, "REPLACED_ACTIVE"),
          "COMPANY_CLOSED": (RESOLVED, "COMPANY_CLOSED"), "MERGED_COMPANY": (RESOLVED, "COMPANY_MERGED"),
          "NOT_FOUND": (RESOLVED, "NOT_FOUND_IN_SEARCH"), "BLOCKED": (UNRESOLVED, "BLOCKED")}
LOW_EVIDENCE_NOTE = "same as the old url provided"
NOTES = [
    f"{LOW_EVIDENCE}: NEEDS_REVIEW rows that could not be confidently mapped and have no suggested domain. "
    f"Their Suggested Domain says '{LOW_EVIDENCE_NOTE}': the only candidate is the URL in 'Old URL Provided', "
    "which could not be proven to belong to the company.",
    "RESOLVED rows carry one domain only, in Final Domain URL. REPLACED_ACTIVE includes 16,509 NEEDS_REVIEW rows "
    "whose suggested domain is now their Final Domain URL (accuracy audit: 74.6% correct, 95% CI 70.3-78.9%).",
    "NOT_FOUND_IN_SEARCH: no website could be found or proven. Audit with a strict Gemini search: 83.7% confirmed "
    "(95% CI 79.8-86.9%).",
    "BLOCKED: the website refused automated checks (access denied, timeout or nearly empty page); not retried.",
]


def classify(row: dict) -> tuple[str, str]:
    status = row["Final Domain Status"]
    if status == "NEEDS_REVIEW":
        return (RESOLVED, "REPLACED_ACTIVE") if (row.get("Suggested Domain") or "").strip() else (UNRESOLVED, LOW_EVIDENCE)
    if status not in DIRECT:
        raise ValueError(f"unexpected Final Domain Status {status!r} (row {row.get('original_company_name')})")
    return DIRECT[status]


def classified_row(row: dict) -> dict:
    super_class, child = classify(row)
    out = dict(row)
    if row["Final Domain Status"] == "NEEDS_REVIEW" and child == "REPLACED_ACTIVE":
        domain = normalize_domain(row["Suggested Domain"]).registered_domain or row["Suggested Domain"].strip()
        out["Final Domain URL"] = f"https://{domain}"
    # One domain per resolved row (Final Domain URL); the Suggested Domain column only marks low-evidence rows.
    out["Suggested Domain"] = LOW_EVIDENCE_NOTE if child == LOW_EVIDENCE else ""
    out.pop("Confidence", None)
    out["Super Class"], out["Final Domain Status"] = super_class, child
    return out


def build(rows: list[dict], columns: list[str]) -> tuple[list[str], list[dict]]:
    """Columns with 'Super Class' before 'Final Domain Status' and no Confidence column; rows sorted by
    ORDER (stable)."""
    columns = [c for c in columns if c != "Confidence"]
    status_at = columns.index("Final Domain Status")
    out_columns = columns[:status_at] + ["Super Class"] + columns[status_at:]
    rank = {key: i for i, key in enumerate(ORDER)}
    out_rows = sorted((classified_row(r) for r in rows), key=lambda r: rank[(r["Super Class"], r["Final Domain Status"])])
    return out_columns, out_rows


def write(columns: list[str], rows: list[dict], out_dir: Path) -> Path:
    with open(out_dir / "final_classified.csv", "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)
    import openpyxl
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("final_classified")
    sheet.append(columns)
    for r in rows:
        sheet.append([excel_cell(r.get(c, "")) for c in columns])
    counts = Counter((r["Super Class"], r["Final Domain Status"]) for r in rows)
    summary = book.create_sheet("summary")
    summary.append(["Super Class", "Final Domain Status", "Rows"])
    for super_class in (RESOLVED, UNRESOLVED):
        for key in ORDER:
            if key[0] == super_class:
                summary.append([super_class, key[1], counts[key]])
        summary.append([super_class, "TOTAL", sum(v for k, v in counts.items() if k[0] == super_class)])
    summary.append(["ALL", "TOTAL", len(rows)])
    summary.append([])
    summary.append(["Notes"])
    for note in NOTES:
        summary.append([note])
    return _save_workbook(book, out_dir / "final_classified.xlsx")


def main() -> None:
    parser = argparse.ArgumentParser(description="RESOLVED / UNRESOLVED classification of the final sheet")
    parser.add_argument("--input", default=str(ROOT / "final_combined" / "final_sorted.csv"))
    parser.add_argument("--out-dir", default=str(ROOT / "final_combined"))
    args = parser.parse_args()
    csv.field_size_limit(1 << 30)
    with open(args.input, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        columns, rows = list(reader.fieldnames or []), list(reader)
    out_columns, out_rows = build(rows, columns)
    saved = write(out_columns, out_rows, Path(args.out_dir))
    counts = Counter((r["Super Class"], r["Final Domain Status"]) for r in out_rows)
    print(f"saved {saved}: {len(out_rows)} rows")
    for key in ORDER:
        print(f"  {key[0]:<10} {key[1]:<34} {counts[key]:>7}")
    for super_class in (RESOLVED, UNRESOLVED):
        print(f"  {super_class:<10} TOTAL {sum(v for k, v in counts.items() if k[0] == super_class):>36}")


if __name__ == "__main__":
    main()
